import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def export_daily_collections_excel(company, payments, target_date):
    """
    Generates an Excel workbook for the daily collection register.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Daily Collections {target_date.strftime('%d-%m-%Y')}"

    # Title styling
    title_font = Font(name="Calibri", size=16, bold=True, color="059669")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid")
    bold_font = Font(name="Calibri", size=11, bold=True)
    border_thin = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    # Company Header
    ws.merge_cells('A1:G1')
    ws['A1'] = company.name
    ws['A1'].font = title_font
    ws['A1'].alignment = Alignment(horizontal="center")

    ws.merge_cells('A2:G2')
    ws['A2'] = f"Daily Collection Register - {target_date.strftime('%d %B %Y')}"
    ws['A2'].font = Font(name="Calibri", size=12, italic=True)
    ws['A2'].alignment = Alignment(horizontal="center")

    headers = [
        "Receipt No", "Time", "Customer Name", "Phone",
        "Loan Account", "Mode", "Amount (₹)", "Collected By"
    ]
    row_num = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    total_amount = 0
    row_num = 5
    for p in payments:
        ws.cell(row=row_num, column=1, value=p.receipt_number).border = border_thin
        ws.cell(row=row_num, column=2, value=p.payment_date.strftime('%H:%M')).border = border_thin
        ws.cell(row=row_num, column=3, value=p.loan.member.name).border = border_thin
        ws.cell(row=row_num, column=4, value=p.loan.member.phone).border = border_thin
        ws.cell(row=row_num, column=5, value=p.loan.loan_account_no).border = border_thin
        ws.cell(row=row_num, column=6, value=p.get_payment_mode_display()).border = border_thin
        
        amt_cell = ws.cell(row=row_num, column=7, value=float(p.amount))
        amt_cell.number_format = '₹#,##0.00'
        amt_cell.border = border_thin
        amt_cell.alignment = Alignment(horizontal="right")

        collector = p.collected_by.get_full_name() or p.collected_by.username
        ws.cell(row=row_num, column=8, value=collector).border = border_thin

        total_amount += float(p.amount)
        row_num += 1

    # Total row
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=6)
    tot_label = ws.cell(row=row_num, column=1, value="TOTAL COLLECTED")
    tot_label.font = bold_font
    tot_label.alignment = Alignment(horizontal="right")
    
    tot_val = ws.cell(row=row_num, column=7, value=total_amount)
    tot_val.font = bold_font
    tot_val.number_format = '₹#,##0.00'
    tot_val.alignment = Alignment(horizontal="right")

    for col in range(1, 9):
        ws.cell(row=row_num, column=col).border = border_thin

    # Auto-adjust column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
